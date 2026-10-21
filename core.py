"""
Appunti Lezioni – motore: configurazione, registrazione, trascrizione, appunti, Word.
Nessuna interfaccia qui: l'app (app.py) comunica con questo modulo tramite callback.
"""

import datetime as dt
import json
import os
import queue
import re
import sys
import threading
import time
import traceback
import uuid
from pathlib import Path

import numpy as np

os.environ.setdefault("HF_HUB_DISABLE_SYMLINKS_WARNING", "1")

APP_NOME = "Appunti Lezioni"
APP_VERSIONE = "3.2.0"
APP_DIR = Path(__file__).resolve().parent


def cartella_dati():
    base = os.environ.get("APPDATA") or str(Path.home() / ".config")
    p = Path(base) / "AppuntiLezioni"
    p.mkdir(parents=True, exist_ok=True)
    return p


DATI = cartella_dati()
CONFIG_PATH = DATI / "config.json"
LOG_PATH = DATI / "appunti_lezioni.log"
SESSIONI_DIR = DATI / "sessioni"
SESSIONI_DIR.mkdir(exist_ok=True)

MODELLI = {"veloce": "base", "standard": "small", "preciso": "medium"}

DEFAULT_CONFIG = {
    "chiave_api_anthropic": "",
    "cartella_appunti": "",
    "qualita_trascrizione": "standard",
    "lingua": "it",
    "modello_claude": "claude-sonnet-5-5",
    "minuti_silenzio_stop_automatico": 8,
    "secondi_per_blocco": 45,
    "soglia_silenzio": 0.002,
    "apri_appunti_al_termine": True,
    "solo_browser": True,
    "contenuti_solo_materiali": "segna",
    "repo_aggiornamenti": "matteopettinato88-dot/appunti-lezioni",
    "ultimo_controllo_aggiornamenti": "",
    "avvio_con_windows": False,
    "corsi": [],
    "ultimo_corso": "",
    "stile_appunti": (
        "Appunti chiari e completi per lo studio universitario:\n"
        "- un titolo con l'argomento principale della lezione\n"
        "- una sezione con sottotitolo per ogni argomento trattato, nell'ordine della lezione\n"
        "- elenchi puntati sintetici, ma senza perdere definizioni, esempi e passaggi logici\n"
        "- definizioni e concetti chiave in grassetto\n"
        "- alla fine una sezione 'Riepilogo' con 5-8 punti\n"
        "- e una sezione 'Possibili domande d'esame' con 3-6 domande"
    ),
}

SYSTEM_PROMPT = """Sei un assistente che trasforma la trascrizione automatica di una lezione universitaria in appunti di studio in italiano.

La trascrizione è stata prodotta da un sistema di riconoscimento vocale: contiene errori (parole sentite male, termini tecnici storpiati, nomi propri sbagliati, frasi spezzate tra un blocco e l'altro), ripetizioni e intercalari tipici del parlato. I segnaposto [min N] indicano il minuto approssimativo della lezione: servono solo a orientarti, non riportarli negli appunti.

Il tuo compito:
1. Ricostruisci il senso corretto usando il contesto della materia: correggi termini tecnici, nomi di autori, leggi, articoli, formule e date.
2. Non inventare contenuti che il docente non ha detto. Se un passaggio è incomprensibile o una correzione è incerta, segnalalo con [?]. Numeri e date sentiti dal riconoscimento vocale possono essere sbagliati: se non sono coerenti col contesto, segnalali con [?].
3. Elimina intercalari, ripetizioni, saluti, inviti a iscriversi al canale, pubblicità e comunicazioni organizzative, ma conserva definizioni, esempi, elenchi, numeri e riferimenti normativi o bibliografici.
4. Scrivi gli appunti seguendo le istruzioni di stile dell'utente.

Se sono allegati materiali (slide, dispense, documenti, foto):
- La lezione resta la fonte principale: gli appunti riportano ciò che il docente ha spiegato.
- Usa i materiali per scrivere correttamente nomi, termini tecnici, date, formule e definizioni, e per dare agli appunti la stessa struttura e gli stessi titoli dei materiali quando corrispondono alla lezione.
- Il minuto in cui un materiale è stato allegato indica più o meno da quando il docente ne parla.
- Se il docente e i materiali dicono cose diverse, riporta la versione del docente e segnala la discordanza con [?] indicando cosa dicono i materiali.
- Per i contenuti presenti solo nei materiali segui la regola indicata nel messaggio.

Se sono elencate lezioni precedenti del corso, quando un concetto è già stato trattato in una di esse aggiungi un breve rimando tra parentesi, ad esempio (vedi Lezione 2 – patto dichiarato). Usa solo le lezioni elencate.

Formato di output (verrà convertito automaticamente in un documento Word). Usa solo:
# Titolo (una sola volta, all'inizio)
## Sezione
### Sottosezione
- elenco puntato (rientra di 2 spazi per i sottopunti)
1. elenco numerato
> nota o concetto da ricordare
**grassetto**, *corsivo* e __sottolineato__
Le istruzioni di stile dell'utente decidono la struttura: titoli, sezioni ed elenchi sono strumenti da usare solo se lo stile li prevede.
Niente tabelle, niente blocchi di codice, niente testo prima del titolo.

Dopo gli appunti aggiungi, per il glossario del corso, un blocco con le 5-15 definizioni più importanti della lezione, in questo formato esatto:
<glossario>
[{"termine": "...", "definizione": "una o due frasi"}]
</glossario>"""


# ----------------------------------------------------------------------------
# Registro e configurazione
# ----------------------------------------------------------------------------

_log_lock = threading.Lock()


def log(msg):
    riga = f"[{dt.datetime.now():%Y-%m-%d %H:%M:%S}] {msg}"
    if sys.stdout is not None:
        try:
            print(riga, flush=True)
        except Exception:
            pass
    with _log_lock:
        try:
            with open(LOG_PATH, "a", encoding="utf-8") as f:
                f.write(riga + "\n")
        except OSError:
            pass


def estrai_chiave(testo):
    """Accetta anche un incolla 'sporco' (es. l'intero comando curl) e tiene solo la chiave."""
    testo = (testo or "").strip()
    m = re.search(r"sk-ant-[A-Za-z0-9_\-]{20,}", testo)
    return m.group(0) if m else testo


def load_config():
    cfg = dict(DEFAULT_CONFIG)
    if CONFIG_PATH.exists():
        try:
            with open(CONFIG_PATH, encoding="utf-8-sig") as f:
                cfg.update(json.load(f))
        except Exception:
            log("config.json illeggibile, uso i valori predefiniti:\n" + traceback.format_exc())
    else:
        _migra_vecchia_config(cfg)
        save_config(cfg)
    cfg["chiave_api_anthropic"] = estrai_chiave(cfg.get("chiave_api_anthropic", ""))
    if not cfg.get("repo_aggiornamenti"):
        cfg["repo_aggiornamenti"] = DEFAULT_CONFIG["repo_aggiornamenti"]
    if cfg.get("qualita_trascrizione") not in MODELLI:
        cfg["qualita_trascrizione"] = "standard"
    return cfg


def _migra_vecchia_config(cfg):
    """Importa la configurazione della prima versione (config.json nella cartella del programma)."""
    vecchia = APP_DIR / "config.json"
    if not vecchia.exists():
        return
    try:
        txt = vecchia.read_text(encoding="utf-8-sig")
        try:
            v = json.loads(txt)
        except Exception:
            v = {"chiave_api_anthropic": estrai_chiave(txt)}
        for k in ("chiave_api_anthropic", "cartella_appunti", "stile_appunti", "ultimo_corso",
                  "minuti_silenzio_stop_automatico", "apri_appunti_al_termine"):
            if v.get(k) not in (None, ""):
                cfg[k] = v[k]
        inv = {m: q for q, m in MODELLI.items()}
        if v.get("modello_whisper") in inv:
            cfg["qualita_trascrizione"] = inv[v["modello_whisper"]]
        if cfg.get("ultimo_corso"):
            cfg["corsi"] = [cfg["ultimo_corso"]]
        log("Configurazione importata dalla versione precedente.")
    except Exception:
        log("Migrazione configurazione fallita:\n" + traceback.format_exc())


def save_config(cfg):
    tmp = CONFIG_PATH.with_suffix(".tmp")
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(cfg, f, indent=2, ensure_ascii=False)
    os.replace(tmp, CONFIG_PATH)


def cartella_documenti():
    if sys.platform == "win32":
        try:
            import ctypes
            buf = ctypes.create_unicode_buffer(1024)
            if ctypes.windll.shell32.SHGetFolderPathW(None, 5, None, 0, buf) == 0 and buf.value:
                return Path(buf.value)
        except Exception:
            pass
    return Path.home() / "Documents"


def cartella_appunti(cfg):
    c = (cfg.get("cartella_appunti") or "").strip()
    return Path(os.path.expandvars(c)).expanduser() if c else cartella_documenti() / "Appunti lezioni"


def pulisci_nome(nome):
    nome = re.sub(r'[\\/:*?"<>|]+', "-", nome or "").strip(" .-")
    nome = re.sub(r"\s+", " ", nome)
    return nome or "Lezione"


def normalizza_corso(nome, corsi):
    """Riusa il nome di un corso esistente se differisce solo per maiuscole/spazi."""
    pulito = pulisci_nome(nome)
    for c in corsi:
        if c.casefold() == pulito.casefold():
            return c
    return pulito[:1].upper() + pulito[1:]


def elenco_corsi(cfg):
    corsi = list(cfg.get("corsi") or [])
    base = cartella_appunti(cfg)
    if base.exists():
        for d in sorted(base.iterdir()):
            if d.is_dir() and not any(c.casefold() == d.name.casefold() for c in corsi):
                corsi.append(d.name)
    return corsi


def appunti_recenti(cfg, n=30):
    base = cartella_appunti(cfg)
    if not base.exists():
        return []
    files = [p for p in base.glob("*/*.docx") if not p.name.startswith("~$")]
    return sorted(files, key=lambda p: p.stat().st_mtime, reverse=True)[:n]


# ----------------------------------------------------------------------------
# Audio
# ----------------------------------------------------------------------------

def a_16k(x, sr):
    """Porta l'audio mono a 16 kHz (il formato di Whisper) con un filtro passa-basso e interpolazione."""
    x = np.asarray(x, dtype=np.float32)
    if sr == 16000 or len(x) == 0:
        return x
    if sr > 16000:
        fc = 7600.0 / sr
        n = np.arange(-50, 51)
        h = 2 * fc * np.sinc(2 * fc * n) * np.hamming(len(n))
        x = np.convolve(x, (h / h.sum()).astype(np.float32), mode="same")
    durata = len(x) / sr
    t_nuovo = np.arange(int(durata * 16000)) / 16000.0
    return np.interp(t_nuovo, np.arange(len(x)) / sr, x).astype(np.float32)


def leggi_file_audio(path):
    import av

    pezzi = []
    with av.open(str(path)) as cont:
        res = av.AudioResampler(format="s16", layout="mono", rate=16000)
        for frame in cont.decode(audio=0):
            for f in res.resample(frame):
                pezzi.append(f.to_ndarray().reshape(-1))
        for f in res.resample(None):
            pezzi.append(f.to_ndarray().reshape(-1))
    return (np.concatenate(pezzi).astype(np.float32) / 32768.0) if pezzi else np.zeros(0, np.float32)


def punto_di_taglio(data, sr, finestra_ricerca_s=10.0, passo_s=0.1):
    """Il punto più silenzioso negli ultimi secondi, per non spezzare le parole."""
    n = len(data)
    inizio = max(0, n - int(finestra_ricerca_s * sr))
    passo = max(1, int(passo_s * sr))
    zona = data[inizio:n]
    k = len(zona) // passo
    if k < 2:
        return n
    energie = (zona[: k * passo].reshape(k, passo) ** 2).mean(axis=1)
    return inizio + int(np.argmin(energie)) * passo + passo // 2


def trova_loopback(pa, pyaudio):
    try:
        wasapi = pa.get_host_api_info_by_type(pyaudio.paWASAPI)
    except OSError:
        raise RuntimeError("Questo programma funziona su Windows 10/11.")
    uscita = pa.get_device_info_by_index(wasapi["defaultOutputDevice"])
    if uscita.get("isLoopbackDevice"):
        return uscita
    for lb in pa.get_loopback_device_info_generator():
        if uscita["name"] in lb["name"]:
            return lb
    raise RuntimeError(f"Non riesco a registrare dall'uscita audio '{uscita['name']}'.")


class Registratore:
    """Registra ciò che esce dalle casse/cuffie e consegna blocchi di audio mono."""

    def __init__(self, cfg, on_blocco):
        import pyaudiowpatch as pyaudio

        self.pyaudio = pyaudio
        self.pa = pyaudio.PyAudio()
        self.dispositivo = trova_loopback(self.pa, pyaudio)
        self.sr = int(self.dispositivo["defaultSampleRate"])
        self.canali = max(1, int(self.dispositivo["maxInputChannels"]))
        self.on_blocco = on_blocco
        self.secondi_blocco = float(cfg["secondi_per_blocco"])
        self.soglia = float(cfg["soglia_silenzio"])
        self.lock = threading.Lock()
        self.buffer = []
        self.campioni = 0
        self.campioni_totali = 0      # tutto l'audio registrato (esclusa la pausa)
        self.consegnati = 0           # campioni già consegnati in blocchi (per i segnaposto [min N])
        self.livello = 0.0
        self.ultimo_suono = time.time()
        self.ultimo_audio = time.time()
        self.in_pausa = False
        self.attivo = False
        self.stream = None
        self.thread = None
        log(f"Registro da: {self.dispositivo['name']} ({self.sr} Hz, {self.canali} canali)")

    def _callback(self, in_data, frame_count, time_info, status):
        if not self.in_pausa and in_data:
            a = np.frombuffer(in_data, dtype=np.int16).astype(np.float32) / 32768.0
            if self.canali > 1:
                a = a[: len(a) // self.canali * self.canali].reshape(-1, self.canali).mean(axis=1)
            if len(a):
                rms = float(np.sqrt(np.mean(a ** 2)))
                self.livello = rms
                self.ultimo_audio = time.time()
                if rms > self.soglia:
                    self.ultimo_suono = time.time()
                with self.lock:
                    self.buffer.append(a)
                    self.campioni += len(a)
                    self.campioni_totali += len(a)
        return (None, self.pyaudio.paContinue)

    def secondi_registrati(self):
        return self.campioni_totali / self.sr

    def livello_attuale(self):
        # Con WASAPI, se nulla suona non arrivano dati: il livello va a zero.
        return 0.0 if time.time() - self.ultimo_audio > 0.5 or self.in_pausa else self.livello

    def avvia(self):
        self.stream = self.pa.open(
            format=self.pyaudio.paInt16, channels=self.canali, rate=self.sr, input=True,
            input_device_index=self.dispositivo["index"], frames_per_buffer=2048,
            stream_callback=self._callback,
        )
        self.stream.start_stream()
        self.attivo = True
        self.ultimo_suono = time.time()
        self.thread = threading.Thread(target=self._ciclo, daemon=True)
        self.thread.start()

    def _ciclo(self):
        while self.attivo:
            time.sleep(1)
            if self.campioni >= self.secondi_blocco * self.sr:
                self._consegna(finale=False)

    def _consegna(self, finale):
        with self.lock:
            dati = np.concatenate(self.buffer) if self.buffer else np.zeros(0, dtype=np.float32)
            self.buffer, self.campioni = [], 0
        if not finale and len(dati) > 12 * self.sr:
            taglio = punto_di_taglio(dati, self.sr)
            resto = dati[taglio:]
            dati = dati[:taglio]
            with self.lock:
                self.buffer.insert(0, resto)
                self.campioni += len(resto)
        inizio = self.consegnati / self.sr
        self.consegnati += len(dati)
        if len(dati) == 0 or float(np.max(np.abs(dati))) < self.soglia:
            return
        self.on_blocco(dati, self.sr, inizio)

    def secondi_di_silenzio(self):
        return time.time() - self.ultimo_suono

    def pausa(self, valore):
        self.in_pausa = valore
        if not valore:
            self.ultimo_suono = time.time()

    def rinvia_stop(self):
        self.ultimo_suono = time.time()

    def ferma(self):
        self.attivo = False
        if self.thread:
            self.thread.join(timeout=5)
        try:
            if self.stream:
                self.stream.stop_stream()
                self.stream.close()
        except Exception:
            pass
        try:
            self.pa.terminate()
        except Exception:
            pass
        self._consegna(finale=True)


# ----------------------------------------------------------------------------
# Trascrizione
# ----------------------------------------------------------------------------

class Trascrittore:
    def __init__(self, qualita, lingua="it", on_stato=None):
        from faster_whisper import WhisperModel

        nome = MODELLI.get(qualita, "small")
        threads = max(1, (os.cpu_count() or 2) - 1)
        t = time.time()
        try:
            self.model = WhisperModel(nome, device="cpu", compute_type="int8",
                                      cpu_threads=threads, local_files_only=True)
        except Exception:
            if on_stato:
                on_stato(f"Scarico il modello di trascrizione '{qualita}' (solo la prima volta)…")
            log(f"Scaricamento modello Whisper '{nome}'…")
            self.model = WhisperModel(nome, device="cpu", compute_type="int8", cpu_threads=threads)
        self.qualita = qualita
        self.lingua = lingua
        log(f"Modello Whisper '{nome}' pronto in {time.time() - t:.0f} s.")

    def trascrivi(self, audio, contesto=""):
        if not isinstance(audio, np.ndarray):
            audio = leggi_file_audio(audio)
        segmenti, _ = self.model.transcribe(
            audio, language=self.lingua, vad_filter=True, beam_size=5,
            initial_prompt=contesto or None,
        )
        return " ".join(s.text.strip() for s in segmenti).strip()


# ----------------------------------------------------------------------------
# Sessione di registrazione con salvataggio continuo
# ----------------------------------------------------------------------------

class Sessione:
    """
    Una lezione in corso. Ogni blocco trascritto viene subito aggiunto a un file su disco
    (con il segnaposto [min N]), così se il PC si spegne la trascrizione non va persa.
    I termini chiave dei materiali (self.termini) guidano Whisper e si possono aggiornare dal vivo.
    """

    def __init__(self, cfg, corso, titolo, trascrittore, on_testo=None, lezione_id=None, termini=None):
        self.cfg = cfg
        self.corso = corso
        self.titolo = titolo
        self.trascrittore = trascrittore
        self.on_testo = on_testo
        self.termini = list(termini or [])
        self.inizio = dt.datetime.now()
        self.id = lezione_id or f"{self.inizio:%Y%m%d-%H%M%S}-{uuid.uuid4().hex[:6]}"
        self.file_testo = SESSIONI_DIR / f"{self.id}.txt"
        self.file_meta = SESSIONI_DIR / f"{self.id}.json"
        self.testi = []
        self.coda = queue.Queue()
        self.n = 0
        self.secondi_audio = 0.0
        self.secondi_calcolo = 0.0
        self.registratore = Registratore(cfg, self._nuovo_blocco)
        self.salva_meta()
        self.file_testo.write_text("", encoding="utf-8")
        self.worker = threading.Thread(target=self._lavora, daemon=True)
        self.worker.start()

    def salva_meta(self):
        self.file_meta.write_text(json.dumps({"corso": self.corso, "titolo": self.titolo, "lezione_id": self.id,
                                              "inizio": self.inizio.isoformat()}, ensure_ascii=False),
                                  encoding="utf-8")

    def minuto_attuale(self):
        return int(self.registratore.secondi_registrati() // 60)

    def _nuovo_blocco(self, dati, sr, inizio_s=0.0):
        self.n += 1
        self.coda.put((self.n, a_16k(dati, sr), len(dati) / sr, inizio_s))

    def _lavora(self):
        import materiali
        while True:
            voce = self.coda.get()
            if voce is None:
                break
            n, audio, durata, inizio_s = voce
            try:
                contesto = " ".join(t for _, t in self.testi)
                prompt = materiali.prompt_whisper(self.corso, self.titolo, self.termini, contesto)
                t = time.time()
                testo = self.trascrittore.trascrivi(audio, prompt)
                calcolo = time.time() - t
                self.secondi_audio += durata
                self.secondi_calcolo += calcolo
                if testo:
                    marcatore = f"[min {int(inizio_s // 60)}]"
                    self.testi.append((marcatore, testo))
                    with open(self.file_testo, "a", encoding="utf-8") as f:
                        f.write(f"{marcatore} {testo}\n\n")
                        f.flush()
                        os.fsync(f.fileno())
                    if self.on_testo:
                        self.on_testo(testo)
                log(f"Blocco {n}: {durata:.0f} s di audio trascritti in {calcolo:.0f} s, {len(testo.split())} parole")
            except Exception:
                log("Errore di trascrizione:\n" + traceback.format_exc())

    def in_ritardo(self):
        return self.coda.qsize()

    def testo_completo(self):
        return "\n\n".join(f"{m} {t}" for m, t in self.testi)

    def termina(self):
        self.registratore.ferma()
        self.coda.put(None)
        self.worker.join()
        return self.testo_completo()

    def chiudi(self):
        """Da chiamare quando la trascrizione è al sicuro nella cartella del corso."""
        elimina_sessione(self.id)


def sessioni_incompiute():
    """Registrazioni interrotte (PC spento, programma chiuso) rimaste in %APPDATA%."""
    risultati = []
    for meta in sorted(SESSIONI_DIR.glob("*.json")):
        testo_p = meta.with_suffix(".txt")
        try:
            m = json.loads(meta.read_text(encoding="utf-8"))
            testo = testo_p.read_text(encoding="utf-8") if testo_p.exists() else ""
        except Exception:
            continue
        risultati.append({"id": meta.stem, "corso": m.get("corso", "Lezione"), "titolo": m.get("titolo", ""),
                          "lezione_id": m.get("lezione_id"),
                          "inizio": dt.datetime.fromisoformat(m["inizio"]),
                          "testo": testo, "parole": len(testo.split())})
    return risultati


def elimina_sessione(id_):
    for ext in (".json", ".txt"):
        try:
            (SESSIONI_DIR / f"{id_}{ext}").unlink()
        except OSError:
            pass


# ----------------------------------------------------------------------------
# Appunti con Claude e Word
# ----------------------------------------------------------------------------

def verifica_chiave(chiave):
    """Restituisce (ok, messaggio)."""
    import anthropic

    chiave = estrai_chiave(chiave)
    if not chiave.startswith("sk-ant-"):
        return False, "La chiave deve iniziare con sk-ant-"
    try:
        client = anthropic.Anthropic(api_key=chiave, max_retries=1, timeout=20)
        client.models.list(limit=1)
        return True, "Chiave valida"
    except anthropic.AuthenticationError:
        return False, "Chiave non valida o disattivata"
    except anthropic.PermissionDeniedError:
        return False, "La chiave non ha i permessi necessari"
    except anthropic.APIConnectionError:
        return False, "Nessuna connessione a Internet"
    except Exception as e:
        return False, f"Errore: {e}"


def messaggio_errore(e):
    """Traduce gli errori più comuni in un messaggio comprensibile."""
    try:
        import anthropic
        if isinstance(e, anthropic.AuthenticationError):
            return "La chiave API non è valida o è scaduta. Controllala nella scheda Impostazioni."
        if isinstance(e, anthropic.PermissionDeniedError):
            return "La chiave API non ha i permessi necessari."
        if isinstance(e, anthropic.RateLimitError):
            return "Troppe richieste in poco tempo: riprova tra qualche minuto."
        if isinstance(e, anthropic.APIConnectionError):
            return "Nessuna connessione a Internet (o il servizio non risponde)."
        if isinstance(e, anthropic.BadRequestError) and "credit" in str(e).lower():
            return "Il credito API è esaurito: ricaricalo su platform.claude.com → Fatturazione."
        if isinstance(e, anthropic.InternalServerError):
            return "Il servizio di Claude è momentaneamente sovraccarico: riprova tra poco."
    except Exception:
        pass
    return str(e)


def crea_appunti(trascrizione, corso, cfg, titolo="", data_lezione=None, materiali_lezione=(), precedenti=()):
    """
    Chiede a Claude gli appunti. Restituisce (markdown, glossario) dove glossario è una lista
    di {"termine", "definizione"} estratti dalla lezione (servono alla dispensa del corso).
    """
    import anthropic
    import materiali

    chiave = estrai_chiave(cfg.get("chiave_api_anthropic") or os.environ.get("ANTHROPIC_API_KEY") or "")
    if not chiave:
        raise RuntimeError("Manca la chiave API: inseriscila nella scheda Impostazioni.")
    client = anthropic.Anthropic(api_key=chiave, max_retries=3)

    blocchi_mat, elenco_mat = materiali.blocchi_claude(list(materiali_lezione))
    intro = [f"Corso: {corso}"]
    if titolo:
        intro.append(f"Lezione: {titolo}" + (f" (del {data_lezione:%d/%m/%Y})" if data_lezione else ""))
    if precedenti:
        intro.append("Lezioni precedenti del corso, con i loro concetti principali (usale solo per i rimandi):\n"
                     + "\n".join(f"- {p}" for p in precedenti))
    intro.append(f"Istruzioni di stile per gli appunti:\n{cfg['stile_appunti']}")
    if elenco_mat:
        extra = cfg.get("contenuti_solo_materiali", "segna")
        regola = ("Aggiungi i contenuti presenti solo nei materiali e non spiegati dal docente, "
                  "segnandoli con *(dai materiali)*." if extra == "segna" else
                  "Non aggiungere contenuti presenti solo nei materiali: riporta solo ciò che il docente ha detto.")
        intro.append("Materiali allegati a questa lezione:\n" + "\n".join(f"- {r}" for r in elenco_mat)
                     + f"\n\n{regola}")

    contenuto = [{"type": "text", "text": "\n\n".join(intro)}]
    contenuto += blocchi_mat
    contenuto.append({"type": "text", "text":
                      f"Trascrizione automatica della lezione:\n<trascrizione>\n{trascrizione}\n</trascrizione>"})

    messaggi = [{"role": "user", "content": contenuto}]
    risposta = ""
    for tentativo in range(3):  # se la risposta viene tagliata per lunghezza, chiede di continuare
        with client.messages.stream(model=cfg["modello_claude"], max_tokens=32000,
                                    system=SYSTEM_PROMPT, messages=messaggi) as stream:
            msg = stream.get_final_message()
        parte = "".join(b.text for b in msg.content if getattr(b, "type", "") == "text")
        risposta += parte
        if msg.stop_reason != "max_tokens":
            break
        log(f"Appunti tagliati per lunghezza: chiedo il seguito (parte {tentativo + 2}).")
        messaggi = [{"role": "user", "content": contenuto},
                    {"role": "assistant", "content": risposta},
                    {"role": "user", "content": "La risposta si è interrotta. Continua esattamente dal punto in "
                                                "cui ti sei fermato, senza ripetere nulla, fino alla fine "
                                                "(glossario compreso)."}]
    else:
        risposta += "\n\n> [?] Gli appunti potrebbero essere incompleti: la lezione era molto lunga."
    return separa_glossario(risposta.strip())


def separa_glossario(risposta):
    glossario = []
    m = re.search(r"<glossario>\s*(.*?)\s*</glossario>", risposta, flags=re.S)
    if m:
        try:
            dati = json.loads(re.sub(r"^```(?:json)?|```$", "", m.group(1).strip()).strip())
            glossario = [{"termine": str(g["termine"]).strip(), "definizione": str(g["definizione"]).strip()}
                         for g in dati if g.get("termine") and g.get("definizione")]
        except Exception:
            log("Glossario non leggibile nella risposta di Claude.")
        risposta = (risposta[:m.start()] + risposta[m.end():]).strip()
    return risposta, glossario


_INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|__[^_]+__|\*[^*\s][^*]*\*)")


def _aggiungi_testo(par, testo):
    for parte in _INLINE_RE.split(testo):
        if not parte:
            continue
        if parte.startswith("**") and parte.endswith("**") and len(parte) > 4:
            par.add_run(parte[2:-2]).bold = True
        elif parte.startswith("__") and parte.endswith("__") and len(parte) > 4:
            par.add_run(parte[2:-2]).underline = True
        elif parte.startswith("*") and parte.endswith("*") and len(parte) > 2:
            par.add_run(parte[1:-1]).italic = True
        else:
            par.add_run(parte)


def nuovo_documento():
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    normale = doc.styles["Normal"]
    normale.font.name = "Calibri"
    normale.font.size = Pt(11)
    return doc


def riga_info(doc, testo):
    from docx.shared import Pt
    p = doc.add_paragraph()
    r = p.add_run(testo)
    r.italic = True
    r.font.size = Pt(9)


def scrivi_markdown(doc, md, sposta_livelli=0, salta_titolo=False, dopo_titolo=None):
    """
    Scrive il markdown nel documento Word. sposta_livelli abbassa i titoli (per i capitoli della
    dispensa); salta_titolo omette la prima riga '# ' (la dispensa usa un suo titolo di capitolo).
    dopo_titolo(doc) viene chiamato subito dopo il titolo principale (o all'inizio se manca).
    """
    from docx.shared import Cm

    fatto = dopo_titolo is None
    for grezza in md.splitlines():
        riga = grezza.rstrip()
        testo = riga.lstrip()
        if not testo or re.fullmatch(r"[-*_]{3,}", testo):
            continue
        rientro = len(riga) - len(testo)

        m = re.match(r"^(#{1,6})\s+(.*)", testo)
        if m:
            livello = min(len(m.group(1)), 3)
            if livello == 1 and salta_titolo:
                salta_titolo = False
                if not fatto:
                    dopo_titolo(doc)
                    fatto = True
                continue
            liv_word = min(3, livello - 1 + sposta_livelli)
            p = doc.add_heading("", level=liv_word)
            _aggiungi_testo(p, m.group(2).strip())
            if livello == 1 and not fatto:
                dopo_titolo(doc)
                fatto = True
            continue
        if not fatto:
            dopo_titolo(doc)
            fatto = True

        m = re.match(r"^[-*•]\s+(.*)", testo)
        if m:
            liv = min(rientro // 2, 2)
            stile = "List Bullet" if liv == 0 else f"List Bullet {liv + 1}"
            _aggiungi_testo(doc.add_paragraph(style=stile), m.group(1))
            continue
        m = re.match(r"^(\d+[.)])\s+(.*)", testo)
        if m:
            liv = min(rientro // 2, 2)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.9 * (liv + 1))
            p.paragraph_format.first_line_indent = Cm(-0.6)
            p.add_run(m.group(1) + " ")
            _aggiungi_testo(p, m.group(2))
            continue
        m = re.match(r"^>\s?(.*)", testo)
        if m:
            _aggiungi_testo(doc.add_paragraph(style="Intense Quote"), m.group(1))
            continue
        _aggiungi_testo(doc.add_paragraph(), testo)
    if not fatto:
        dopo_titolo(doc)


def markdown_in_word(md, path_docx, info):
    doc = nuovo_documento()
    scrivi_markdown(doc, md, dopo_titolo=lambda d: riga_info(d, info))
    doc.save(str(path_docx))


def percorso_libero(path):
    """Se il file esiste già (o è aperto in Word), aggiunge (2), (3)…"""
    path = Path(path)
    if not path.exists():
        return path
    i = 2
    while True:
        p = path.with_name(f"{path.stem} ({i}){path.suffix}")
        if not p.exists():
            return p
        i += 1
