"""
Appunti Lezioni – materiali allegati alle lezioni (PDF, PowerPoint, Word, testo, foto).
Estrae il testo, i termini chiave per Whisper e prepara i blocchi da inviare a Claude.
"""

import base64
import io
import os
import re
import time
from collections import Counter
from pathlib import Path

from core import log

TIPI = {
    ".pdf": "pdf", ".pptx": "pptx", ".docx": "docx", ".txt": "testo", ".md": "testo",
    ".jpg": "immagine", ".jpeg": "immagine", ".png": "immagine", ".webp": "immagine",
    ".gif": "immagine", ".bmp": "immagine",
}
ICONE = {"pdf": "📄", "pptx": "📊", "docx": "📝", "testo": "📝", "immagine": "🖼"}
VECCHI_FORMATI = {".ppt": "PowerPoint", ".doc": "Word"}

LIMITE_PAGINE_PDF = 90           # oltre, i PDF vengono passati come solo testo
LIMITE_BYTES_DOCUMENTI = 20_000_000
LIMITE_TESTO_MATERIALE = 60_000  # caratteri per singolo materiale
LIMITE_TESTO_TOTALE = 160_000
LIMITE_IMMAGINI = 15


def tipo(path):
    return TIPI.get(Path(path).suffix.lower())


def icona(path):
    return ICONE.get(tipo(path), "📎")


def controlla(path):
    """None se il file è accettato, altrimenti un messaggio per l'utente."""
    p = Path(path)
    ext = p.suffix.lower()
    if ext in VECCHI_FORMATI:
        return (f"«{p.name}» è nel vecchio formato {VECCHI_FORMATI[ext]}. Aprilo e salvalo come "
                f"{'.pptx' if ext == '.ppt' else '.docx'}, poi allegalo di nuovo.")
    if ext not in TIPI:
        return f"«{p.name}»: formato non supportato. Si possono allegare PDF, PowerPoint, Word, testo e foto."
    if not p.exists():
        return f"«{p.name}» non esiste più."
    return None


# ----------------------------------------------------------------------------
# Estrazione del testo
# ----------------------------------------------------------------------------

def pagine_pdf(path):
    try:
        from pypdf import PdfReader
        return len(PdfReader(str(path)).pages)
    except Exception:
        return 0


def testo(path):
    """Testo leggibile del materiale (stringa vuota per le immagini o in caso di errore)."""
    p = Path(path)
    t = tipo(p)
    try:
        if t == "pdf":
            from pypdf import PdfReader
            r = PdfReader(str(p))
            return "\n\n".join(f"[Pagina {i + 1}]\n{(pg.extract_text() or '').strip()}"
                               for i, pg in enumerate(r.pages))
        if t == "pptx":
            from pptx import Presentation
            pres = Presentation(str(p))
            parti = []
            for i, slide in enumerate(pres.slides, 1):
                righe = []
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for par in shape.text_frame.paragraphs:
                            s = "".join(r.text for r in par.runs).strip()
                            if s:
                                righe.append(s)
                    if getattr(shape, "has_table", False) and shape.has_table:
                        for row in shape.table.rows:
                            righe.append(" | ".join(c.text.strip() for c in row.cells))
                if slide.has_notes_slide:
                    note = slide.notes_slide.notes_text_frame.text.strip()
                    if note:
                        righe.append(f"(Note del relatore: {note})")
                parti.append(f"[Slide {i}]\n" + "\n".join(righe))
            return "\n\n".join(parti)
        if t == "docx":
            from docx import Document
            d = Document(str(p))
            righe = [par.text for par in d.paragraphs if par.text.strip()]
            for tab in d.tables:
                for row in tab.rows:
                    righe.append(" | ".join(c.text.strip() for c in row.cells))
            return "\n".join(righe)
        if t == "testo":
            for enc in ("utf-8-sig", "cp1252", "latin-1"):
                try:
                    return p.read_text(encoding=enc)
                except UnicodeDecodeError:
                    continue
    except Exception as e:
        log(f"Testo non estratto da {p.name}: {e}")
    return ""


# ----------------------------------------------------------------------------
# Termini chiave per Whisper
# ----------------------------------------------------------------------------

_COMUNI = set("""
il lo la i gli le un uno una di da in con su per tra fra a e o ma se che chi cui non più meno
del dello della dei degli delle al allo alla ai agli alle dal dallo dalla dai dagli dalle nel nello nella
nei negli nelle sul sullo sulla sui sugli sulle come quando dove perché anche quindi dunque questo questa
questi queste quello quella quelli quelle sono essere stato stata hanno avere viene vengono può possono
deve devono molto poco ogni tutto tutti tutte altro altri altre loro suo sua suoi sue esempio ovvero
cioè inoltre infatti mentre secondo rispetto attraverso sempre spesso ancora proprio parte modo caso
tipo cosa cose fatto fatti livello aspetti aspetto elemento elementi processo processi sistema sistemi
diverse diversi diverso diversa importante importanti generale particolare relativo relativa relazione
the and of to in for on with by from that this these those are is be as at an or not it its into
pagina slide lezione lezioni corso capitolo introduzione conclusioni conclusione domande domanda
""".split())


def termini_chiave(testo_materiale, n=40):
    """
    Nomi propri, sigle, anni e parole straniere/tecniche scritte con la maiuscola: sono le parole
    che Whisper sbaglia più spesso. Le parole italiane comuni non servono (Whisper le conosce già).
    """
    if not testo_materiale:
        return []
    t = re.sub(r"\[(Pagina|Slide) \d+\]", " ", testo_materiale)
    maiuscole, minuscole, sigle, forma = Counter(), Counter(), Counter(), {}
    for m in re.finditer(r"[A-Za-zÀ-ÖØ-öø-ÿ][A-Za-zÀ-ÖØ-öø-ÿ\-]{1,}", t):
        w = m.group(0).strip("-")
        k = w.casefold()
        if len(w) < 2 or k in _COMUNI:
            continue
        if re.fullmatch(r"[A-ZÀ-Ö]{2,7}s?", w) and len(w) >= 2:
            sigle[w] += 1
            continue
        prima = t[max(0, m.start() - 3):m.start()]
        inizio_frase = m.start() == 0 or bool(re.search(r"([.!?:;\n•\-–—]|^)\s*$", prima))
        if w[0].isupper():
            if not inizio_frase and len(w) >= 3:
                maiuscole[k] += 1
                forma.setdefault(k, w)
        else:
            minuscole[k] += 1
    punteggi = Counter()
    for k, v in maiuscole.items():
        if v > minuscole[k]:                 # di solito scritta con la maiuscola: nome o termine proprio
            punteggi[forma[k]] = v * 2
    for s, v in sigle.items():
        punteggi[s] = v * 3
    for m in re.finditer(r"\b(1[6-9]\d\d|20\d\d)\b", t):
        punteggi[m.group(0)] += 1
    return [w for w, _ in punteggi.most_common(n)]


def prompt_whisper(corso, titolo, termini, contesto="", max_car=700):
    base = f"Lezione universitaria di {corso}"
    if titolo:
        base += f": {titolo}"
    base += "."
    if termini:
        lista = ""
        for t in termini:
            if len(base) + len(lista) + len(t) + 14 > max_car - 160:
                break
            lista += (", " if lista else "") + t
        base += f" Termini: {lista}."
    return (base + " " + contesto[-160:]).strip()


# ----------------------------------------------------------------------------
# Blocchi da inviare a Claude
# ----------------------------------------------------------------------------

def _immagine_b64(path, lato=1568):
    from PIL import Image, ImageOps
    img = ImageOps.exif_transpose(Image.open(path)).convert("RGB")
    img.thumbnail((lato, lato))
    buf = io.BytesIO()
    img.save(buf, "JPEG", quality=85)
    return base64.b64encode(buf.getvalue()).decode()


def blocchi_claude(materiali):
    """
    materiali: lista di dict {path, nome, minuto}. Restituisce (blocchi, descrizione).
    I PDF vanno a Claude interi (testo + immagini) finché si resta nei limiti, poi come testo.
    """
    blocchi, righe = [], []
    pagine_usate = bytes_usati = testo_usato = immagini = 0
    for i, m in enumerate(materiali, 1):
        p = Path(m["path"])
        if not p.exists():
            continue
        t = tipo(p)
        quando = f"allegato al minuto {m['minuto']} della lezione" if m.get("minuto") is not None else "allegato"
        etichetta = f"Materiale {i}: «{m['nome']}» ({quando})"
        righe.append(etichetta)
        try:
            if t == "pdf":
                n = pagine_pdf(p)
                dim = p.stat().st_size
                if n and pagine_usate + n <= LIMITE_PAGINE_PDF and bytes_usati + dim <= LIMITE_BYTES_DOCUMENTI:
                    blocchi.append({"type": "text", "text": etichetta})
                    blocchi.append({"type": "document", "title": m["nome"],
                                    "source": {"type": "base64", "media_type": "application/pdf",
                                               "data": base64.b64encode(p.read_bytes()).decode()}})
                    pagine_usate += n
                    bytes_usati += dim
                    continue
            if t == "immagine":
                if immagini < LIMITE_IMMAGINI:
                    blocchi.append({"type": "text", "text": etichetta})
                    blocchi.append({"type": "image", "source": {"type": "base64", "media_type": "image/jpeg",
                                                                "data": _immagine_b64(p)}})
                    immagini += 1
                continue
            contenuto = testo(p)[:LIMITE_TESTO_MATERIALE]
            if not contenuto.strip() or testo_usato >= LIMITE_TESTO_TOTALE:
                continue
            contenuto = contenuto[: LIMITE_TESTO_TOTALE - testo_usato]
            testo_usato += len(contenuto)
            blocchi.append({"type": "text", "text": f"{etichetta}\n<materiale>\n{contenuto}\n</materiale>"})
        except Exception as e:
            log(f"Materiale non utilizzabile {p.name}: {e}")
    return blocchi, righe


# ----------------------------------------------------------------------------
# Proposta automatica: l'ultimo file scaricato
# ----------------------------------------------------------------------------

def cartella_download():
    if os.name == "nt":
        try:
            import ctypes
            from ctypes import wintypes
            from uuid import UUID

            class GUID(ctypes.Structure):
                _fields_ = [("Data1", wintypes.DWORD), ("Data2", wintypes.WORD),
                            ("Data3", wintypes.WORD), ("Data4", wintypes.BYTE * 8)]

            u = UUID("{374DE290-123F-4565-9164-39C4925E467B}")  # FOLDERID_Downloads
            g = GUID(u.fields[0], u.fields[1], u.fields[2], (wintypes.BYTE * 8).from_buffer_copy(u.bytes[8:]))
            p = ctypes.c_wchar_p()
            if ctypes.windll.shell32.SHGetKnownFolderPath(ctypes.byref(g), 0, None, ctypes.byref(p)) == 0:
                percorso = Path(p.value)
                ctypes.windll.ole32.CoTaskMemFree(p)
                return percorso
        except Exception:
            pass
    return Path.home() / "Downloads"


def ultimo_scaricato(gia_usati=(), giorni=3):
    """Il materiale più recente nella cartella Download (ultimi giorni), se non è già stato usato."""
    d = cartella_download()
    if not d.exists():
        return None
    limite = time.time() - giorni * 86400
    usati = {n.casefold() for n in gia_usati}
    candidati = []
    try:
        for p in d.iterdir():
            if p.is_file() and tipo(p) in ("pdf", "pptx", "docx") and p.name.casefold() not in usati:
                mt = p.stat().st_mtime
                if mt >= limite:
                    candidati.append((mt, p))
    except OSError:
        return None
    return max(candidati)[1] if candidati else None
