"""
Appunti Lezioni – corsi, lezioni, materiali e dispensa.

Ogni corso è una cartella dentro «Appunti lezioni»:
    <Corso>/
        corso.json                     elenco lezioni, materiali, glossario
        Lezione 4 – Genitorialità.docx appunti delle singole lezioni
        Dispensa – <Corso>.docx         la dispensa a capitoli (e il PDF, se esportato)
        materiali/<Lezione …>/          copie dei materiali allegati
        trascrizioni/                   trascrizioni grezze
        _dati/<id>.md                   appunti in formato testo (per ricostruire la dispensa)
"""

import datetime as dt
import json
import os
import re
import shutil
import threading
import traceback
from pathlib import Path

import core
import materiali
from core import log

_lock = threading.RLock()


# ----------------------------------------------------------------------------
# Dati del corso
# ----------------------------------------------------------------------------

def cartella(cfg, corso):
    return core.cartella_appunti(cfg) / core.pulisci_nome(corso)


def _file_dati(cfg, corso):
    return cartella(cfg, corso) / "corso.json"


def leggi(cfg, corso):
    with _lock:
        p = _file_dati(cfg, corso)
        if p.exists():
            try:
                return json.loads(p.read_text(encoding="utf-8"))
            except Exception:
                log(f"corso.json illeggibile per {corso}:\n" + traceback.format_exc())
        return {"corso": corso, "lezioni": []}


def scrivi(cfg, corso, dati):
    with _lock:
        c = cartella(cfg, corso)
        c.mkdir(parents=True, exist_ok=True)
        tmp = c / "corso.json.tmp"
        tmp.write_text(json.dumps(dati, indent=2, ensure_ascii=False), encoding="utf-8")
        os.replace(tmp, c / "corso.json")


def modifica_lezione(cfg, corso, lezione_id, fn):
    """Applica fn(lezione) e salva. Restituisce la lezione aggiornata."""
    with _lock:
        dati = leggi(cfg, corso)
        lez = next((l for l in dati["lezioni"] if l["id"] == lezione_id), None)
        if lez is None:
            raise KeyError(f"Lezione {lezione_id} non trovata in {corso}")
        fn(lez)
        scrivi(cfg, corso, dati)
        return lez


def lezione(cfg, corso, lezione_id):
    return next((l for l in leggi(cfg, corso)["lezioni"] if l["id"] == lezione_id), None)


def lezioni(cfg, corso):
    return sorted(leggi(cfg, corso)["lezioni"], key=lambda l: (l.get("numero", 0), l["inizio"]))


def prossimo_numero(cfg, corso):
    return max([l.get("numero", 0) for l in leggi(cfg, corso)["lezioni"]] + [0]) + 1


def titolo_proposto(cfg, corso):
    return f"Lezione {prossimo_numero(cfg, corso)}" if corso.strip() else "Lezione 1"


def corsi_con_lezioni(cfg):
    base = core.cartella_appunti(cfg)
    return sorted(d.name for d in base.iterdir() if (d / "corso.json").exists()) if base.exists() else []


def nuova_lezione(cfg, corso, titolo, lezione_id, inizio):
    titolo = (titolo or "").strip() or titolo_proposto(cfg, corso)
    m = re.match(r"(?i)lezione\s+(\d+)", titolo)
    with _lock:
        dati = leggi(cfg, corso)
        numeri = [l.get("numero", 0) for l in dati["lezioni"]]
        numero = int(m.group(1)) if m and int(m.group(1)) not in numeri else max(numeri + [0]) + 1
        lez = {"id": lezione_id, "numero": numero, "titolo": titolo, "inizio": inizio.isoformat(),
               "stato": "in_corso", "appunti": None, "trascrizione": None, "materiali": [],
               "materiali_modificati": False, "glossario": []}
        dati["lezioni"].append(lez)
        scrivi(cfg, corso, dati)
    return lez


def rinomina_lezione(cfg, corso, lezione_id, nuovo_corso, nuovo_titolo):
    """Cambia titolo e/o corso di una lezione (anche durante la registrazione)."""
    nuovo_titolo = nuovo_titolo.strip() or titolo_proposto(cfg, nuovo_corso)
    if core.pulisci_nome(nuovo_corso).casefold() == core.pulisci_nome(corso).casefold():
        return modifica_lezione(cfg, corso, lezione_id, lambda l: l.update(titolo=nuovo_titolo))
    with _lock:
        vecchio = leggi(cfg, corso)
        lez = next(l for l in vecchio["lezioni"] if l["id"] == lezione_id)
        vecchio["lezioni"].remove(lez)
        nuovo = leggi(cfg, nuovo_corso)
        lez["titolo"] = nuovo_titolo
        lez["numero"] = max([l.get("numero", 0) for l in nuovo["lezioni"]] + [0]) + 1
        for m in lez["materiali"]:  # sposta le copie dei materiali
            src = cartella(cfg, corso) / m["file"]
            dst = core.percorso_libero(_cartella_materiali(cfg, nuovo_corso, lez) / Path(m["file"]).name)
            if src.exists():
                dst.parent.mkdir(parents=True, exist_ok=True)
                shutil.move(str(src), str(dst))
                m["file"] = str(dst.relative_to(cartella(cfg, nuovo_corso)))
        nuovo["lezioni"].append(lez)
        scrivi(cfg, nuovo_corso, nuovo)
        scrivi(cfg, corso, vecchio)
    return lez


def elimina_lezione(cfg, corso, lezione_id):
    with _lock:
        dati = leggi(cfg, corso)
        dati["lezioni"] = [l for l in dati["lezioni"] if l["id"] != lezione_id]
        scrivi(cfg, corso, dati)


# ----------------------------------------------------------------------------
# Materiali
# ----------------------------------------------------------------------------

def _cartella_materiali(cfg, corso, lez):
    return cartella(cfg, corso) / "materiali" / core.pulisci_nome(f"{lez.get('numero', '')} - {lez['titolo']}")


def aggiungi_materiale(cfg, corso, lezione_id, sorgente, minuto=None):
    """Copia il file nella cartella del corso e lo collega alla lezione. Restituisce il materiale."""
    sorgente = Path(sorgente)
    errore = materiali.controlla(sorgente)
    if errore:
        raise ValueError(errore)
    lez = lezione(cfg, corso, lezione_id)
    dest_dir = _cartella_materiali(cfg, corso, lez)
    dest_dir.mkdir(parents=True, exist_ok=True)
    dest = core.percorso_libero(dest_dir / sorgente.name)
    shutil.copy2(sorgente, dest)
    mat = {"file": str(dest.relative_to(cartella(cfg, corso))), "nome": sorgente.name,
           "origine": str(sorgente), "minuto": minuto, "aggiunto": dt.datetime.now().isoformat()}

    def agg(l):
        l["materiali"].append(mat)
        if l["stato"] == "completa":
            l["materiali_modificati"] = True
    modifica_lezione(cfg, corso, lezione_id, agg)
    log(f"Materiale allegato a {corso} / {lez['titolo']}: {sorgente.name}" + (f" (min {minuto})" if minuto is not None else ""))
    return mat


def rimuovi_materiale(cfg, corso, lezione_id, file_rel):
    def rim(l):
        l["materiali"] = [m for m in l["materiali"] if m["file"] != file_rel]
        if l["stato"] == "completa":
            l["materiali_modificati"] = True
    modifica_lezione(cfg, corso, lezione_id, rim)
    try:
        (cartella(cfg, corso) / file_rel).unlink()
    except OSError:
        pass


def percorso_materiale(cfg, corso, mat):
    return cartella(cfg, corso) / mat["file"]


def materiali_per_claude(cfg, corso, lez):
    return [{"path": percorso_materiale(cfg, corso, m), "nome": m["nome"], "minuto": m.get("minuto")}
            for m in lez.get("materiali", [])]


def termini_lezione(cfg, corso, lez, glossario_corso=True):
    """Vocabolario per Whisper: termini dei materiali della lezione + glossario del corso."""
    termini = []
    for m in lez.get("materiali", []):
        termini += materiali.termini_chiave(materiali.testo(percorso_materiale(cfg, corso, m)))
    if glossario_corso:
        for l in leggi(cfg, corso)["lezioni"]:
            termini += [g["termine"] for g in l.get("glossario", [])]
    visti, unici = set(), []
    for t in termini:
        if t.casefold() not in visti:
            visti.add(t.casefold())
            unici.append(t)
    return unici[:60]


def materiali_recenti(cfg, n=8):
    """Gli ultimi materiali usati in tutti i corsi (per ricollegarli con un clic)."""
    tutti = []
    for corso in corsi_con_lezioni(cfg):
        for l in leggi(cfg, corso)["lezioni"]:
            for m in l.get("materiali", []):
                p = percorso_materiale(cfg, corso, m)
                if p.exists():
                    tutti.append((m.get("aggiunto", ""), m["nome"], p))
    visti, risultato = set(), []
    for _, nome, p in sorted(tutti, reverse=True):
        if nome.casefold() not in visti:
            visti.add(nome.casefold())
            risultato.append((nome, p))
        if len(risultato) >= n:
            break
    return risultato


def nomi_materiali_usati(cfg):
    return {m["nome"] for c in corsi_con_lezioni(cfg) for l in leggi(cfg, c)["lezioni"] for m in l.get("materiali", [])}


# ----------------------------------------------------------------------------
# Dalla trascrizione agli appunti
# ----------------------------------------------------------------------------

def salva_trascrizione(cfg, corso, lezione_id, testo):
    lez = lezione(cfg, corso, lezione_id)
    d = cartella(cfg, corso) / "trascrizioni"
    d.mkdir(parents=True, exist_ok=True)
    p = d / f"{core.pulisci_nome(lez['titolo'])} ({lez['id'][:8]}).txt"
    p.write_text(testo, encoding="utf-8")
    rel = str(p.relative_to(cartella(cfg, corso)))
    modifica_lezione(cfg, corso, lezione_id,
                     lambda l: l.update(trascrizione=rel, stato="da_completare" if l["stato"] == "in_corso" else l["stato"],
                                        fine=dt.datetime.now().isoformat()))
    return p


def _precedenti(cfg, corso, lez):
    righe = []
    for l in lezioni(cfg, corso):
        if l["id"] == lez["id"] or l.get("stato") != "completa" or l.get("numero", 0) >= lez.get("numero", 0):
            continue
        termini = ", ".join(g["termine"] for g in l.get("glossario", [])[:12])
        righe.append(titolo_capitolo(l) + (f": {termini}" if termini else ""))
    return righe


def crea_appunti_lezione(cfg, corso, lezione_id):
    """Crea (o rigenera) gli appunti Word della lezione e aggiorna la dispensa. Restituisce il percorso."""
    lez = lezione(cfg, corso, lezione_id)
    if not lez.get("trascrizione"):
        raise RuntimeError("Questa lezione non ha una trascrizione.")
    trascrizione = (cartella(cfg, corso) / lez["trascrizione"]).read_text(encoding="utf-8")
    if not trascrizione.strip():
        raise RuntimeError("La trascrizione di questa lezione è vuota.")
    inizio = dt.datetime.fromisoformat(lez["inizio"])
    md, glossario = core.crea_appunti(trascrizione, corso, cfg, titolo=lez["titolo"], data_lezione=inizio,
                                      materiali_lezione=materiali_per_claude(cfg, corso, lez),
                                      precedenti=_precedenti(cfg, corso, lez))
    c = cartella(cfg, corso)
    (c / "_dati").mkdir(exist_ok=True)
    (c / "_dati" / f"{lez['id']}.md").write_text(md, encoding="utf-8")

    vecchio = c / lez["appunti"] if lez.get("appunti") else None
    destinazione = vecchio or c / f"{core.pulisci_nome(lez['titolo'])}.docx"
    if not vecchio:
        destinazione = core.percorso_libero(destinazione)
    nmat = len(lez.get("materiali", []))
    info = f"{corso} · {lez['titolo']} · {inizio:%d/%m/%Y}" + (f" · {nmat} materiali" if nmat else "")
    try:
        core.markdown_in_word(md, destinazione, info)
    except PermissionError:  # file aperto in Word
        destinazione = core.percorso_libero(destinazione)
        core.markdown_in_word(md, destinazione, info)
    rel = str(destinazione.relative_to(c))
    modifica_lezione(cfg, corso, lezione_id, lambda l: l.update(
        appunti=rel, stato="completa", glossario=glossario, materiali_modificati=False,
        creati=dt.datetime.now().isoformat()))
    log(f"Appunti salvati in {destinazione}")
    try:
        aggiorna_dispensa(cfg, corso)
    except Exception:
        log("Dispensa non aggiornata:\n" + traceback.format_exc())
    return destinazione


def importa_sessioni_interrotte(cfg):
    """
    Le registrazioni rimaste in %APPDATA% (programma chiuso o PC spento durante la lezione)
    diventano lezioni «da completare» nel loro corso. Restituisce quante ne ha importate.
    """
    n = 0
    for s in core.sessioni_incompiute():
        try:
            corso = s["corso"]
            lid = s.get("lezione_id") or s["id"]
            if not lezione(cfg, corso, lid):
                nuova_lezione(cfg, corso, s.get("titolo") or "", lid, s["inizio"])
            if s["testo"].strip():
                salva_trascrizione(cfg, corso, lid, s["testo"])
            else:
                elimina_lezione(cfg, corso, lid)
            core.elimina_sessione(s["id"])
            n += 1
        except Exception:
            log("Importazione sessione fallita:\n" + traceback.format_exc())
    return n


def lezioni_da_completare(cfg, escludi_id=None):
    out = []
    for corso in corsi_con_lezioni(cfg):
        for l in lezioni(cfg, corso):
            if l["id"] != escludi_id and l["stato"] in ("da_completare", "in_corso"):
                out.append((corso, l))
    return out


# ----------------------------------------------------------------------------
# Dispensa a capitoli
# ----------------------------------------------------------------------------

def percorso_dispensa(cfg, corso):
    return cartella(cfg, corso) / f"Dispensa – {core.pulisci_nome(corso)}.docx"


def aggiorna_dispensa(cfg, corso):
    """Ricostruisce la dispensa: indice, un capitolo per lezione completata, glossario con rimandi."""
    from docx.enum.text import WD_BREAK

    c = cartella(cfg, corso)
    complete = [l for l in lezioni(cfg, corso)
                if l["stato"] == "completa" and (c / "_dati" / f"{l['id']}.md").exists()]
    if not complete:
        return None
    doc = core.nuovo_documento()
    doc.add_heading(f"{corso}", level=0)
    core.riga_info(doc, f"Dispensa del corso · {len(complete)} lezioni · aggiornata il "
                        f"{dt.datetime.now():%d/%m/%Y %H:%M}")
    doc.add_heading("Indice", level=1)
    for l in complete:
        doc.add_paragraph(titolo_capitolo(l), style="List Number")
    doc.add_paragraph().add_run("Glossario dei termini del corso (in fondo).").italic = True

    for l in complete:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        doc.add_heading(titolo_capitolo(l), level=1)
        inizio = dt.datetime.fromisoformat(l["inizio"])
        mats = ", ".join(m["nome"] for m in l.get("materiali", []))
        core.riga_info(doc, f"{inizio:%d/%m/%Y}" + (f" · materiali: {mats}" if mats else ""))
        md = (c / "_dati" / f"{l['id']}.md").read_text(encoding="utf-8")
        core.scrivi_markdown(doc, md, sposta_livelli=1, salta_titolo=True)

    voci = {}
    for l in complete:
        for g in l.get("glossario", []):
            k = g["termine"].casefold()
            if k not in voci:
                voci[k] = (g["termine"], g["definizione"], [l["numero"]])
            elif l["numero"] not in voci[k][2]:
                voci[k][2].append(l["numero"])
    if voci:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        doc.add_heading("Glossario", level=1)
        for _, (termine, definizione, numeri) in sorted(voci.items()):
            p = doc.add_paragraph()
            p.add_run(termine).bold = True
            p.add_run(f": {definizione} ")
            r = p.add_run("(Lezione " + ", ".join(str(n) for n in sorted(numeri)) + ")")
            r.italic = True

    dest = percorso_dispensa(cfg, corso)
    try:
        doc.save(str(dest))
    except PermissionError:
        dest = core.percorso_libero(dest)
        doc.save(str(dest))
    dati = leggi(cfg, corso)
    dati["dispensa_aggiornata"] = dt.datetime.now().isoformat()
    scrivi(cfg, corso, dati)
    log(f"Dispensa aggiornata: {dest}")
    return dest


def _senza_numero(titolo):
    """«Lezione 4 – Genitorialità» → «Genitorialità» (il numero lo mette già la dispensa)."""
    return re.sub(r"(?i)^lezione\s+\d+\s*[-–—:.]?\s*", "", titolo).strip()


def titolo_capitolo(l):
    resto = _senza_numero(l["titolo"])
    return f"Lezione {l['numero']}" + (f" – {resto}" if resto else "")


def esporta_pdf(docx_path):
    """Esporta in PDF usando Microsoft Word (se installato). Restituisce il percorso del PDF."""
    import comtypes
    import comtypes.client

    docx_path = Path(docx_path).resolve()
    pdf = docx_path.with_suffix(".pdf")
    comtypes.CoInitialize()
    word = comtypes.client.CreateObject("Word.Application")
    try:
        word.Visible = False
        doc = word.Documents.Open(str(docx_path), ReadOnly=True)
        doc.ExportAsFixedFormat(str(pdf), 17)  # 17 = PDF
        doc.Close(False)
    finally:
        word.Quit()
    return pdf
